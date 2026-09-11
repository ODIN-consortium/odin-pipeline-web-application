"""Tests for /api/export/excel (GET) and /api/export/excel/import (POST)."""

import io
import sqlite3

import openpyxl
from fastapi.testclient import TestClient

_XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _make_xlsx(sheets: dict[str, list[list]]) -> bytes:
    """Build an in-memory xlsx with the given sheets (name → list of row lists)."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _upload(
    client: TestClient,
    xlsx_bytes: bytes,
    mode: str = "merge",
    filename: str = "test.xlsx",
    import_valid_rows: bool = False,
) -> object:
    """Upload a workbook.

    `import_valid_rows` mirrors the API default: an upload containing any unimportable row is
    rejected with 422 and nothing is written unless the caller explicitly opts in. Tests that
    exercise per-row skipping therefore pass True; tests of the rejection contract do not.
    """
    return client.post(
        "/api/export/excel/import",
        files={"file": (filename, xlsx_bytes, _XLSX_CONTENT_TYPE)},
        params={"mode": mode, "import_valid_rows": import_valid_rows},
    )


def _site(client: TestClient, **kw) -> dict:
    r = client.post("/api/sites", json={
        "country": "Norway", "country_code": "NO", "city_code": "BGO", "site": "River", **kw,
    })
    assert r.status_code == 201, r.text
    return r.json()


# ── GET /export/excel ─────────────────────────────────────────────────────────


def test_export_excel_returns_200(client: TestClient) -> None:
    r = client.get("/api/export/excel")
    assert r.status_code == 200


def test_export_excel_content_type_is_xlsx(client: TestClient) -> None:
    r = client.get("/api/export/excel")
    assert "spreadsheetml.sheet" in r.headers["content-type"]


def test_export_excel_has_all_expected_sheets(client: TestClient) -> None:
    r = client.get("/api/export/excel")
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert set(wb.sheetnames) == {"sites", "samples", "nanopore", "biomeme"}


def test_export_excel_empty_db_has_header_only_rows(client: TestClient) -> None:
    r = client.get("/api/export/excel")
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb["sites"].max_row == 1
    assert wb["samples"].max_row == 1
    assert wb["nanopore"].max_row == 1
    assert wb["biomeme"].max_row == 1


def test_export_excel_sites_first_row_is_header(client: TestClient) -> None:
    r = client.get("/api/export/excel")
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    ws = wb["sites"]
    header = [cell.value for cell in ws[1]]
    assert "site_code" in header
    assert "country" in header


def test_export_excel_includes_site_data(client: TestClient) -> None:
    _site(client)
    r = client.get("/api/export/excel")
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    ws = wb["sites"]
    rows = list(ws.iter_rows(values_only=True))
    site_codes = [row[0] for row in rows[1:] if row[0]]
    assert "NOBGORiver" in site_codes


def test_export_excel_includes_sample_data(client: TestClient) -> None:
    site = _site(client)
    client.post("/api/samples", json={
        "site_id": site["id"], "sampling_date": "20240601", "sample_type": "water",
    })
    r = client.get("/api/export/excel")
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb["samples"].max_row == 2  # header + 1 data row


def test_export_excel_content_disposition_filename(client: TestClient) -> None:
    r = client.get("/api/export/excel")
    assert "odin_metadata.xlsx" in r.headers.get("content-disposition", "")


# ── POST /export/excel/import — validation ────────────────────────────────────


def test_import_rejects_non_xlsx_extension(client: TestClient) -> None:
    r = _upload(client, b"fake data", filename="data.csv")
    assert r.status_code == 400


def test_import_rejects_invalid_xlsx_content(client: TestClient) -> None:
    r = _upload(client, b"this is not an xlsx file")
    assert r.status_code == 400


def test_import_empty_workbook_returns_zero_summary(client: TestClient) -> None:
    xlsx = _make_xlsx({"sites": [["site_code", "country"]]})
    r = _upload(client, xlsx)
    assert r.status_code == 200
    s = r.json()["summary"]
    assert s["sites"]["created"] == 0
    assert s["sites"]["updated"] == 0


def test_import_returns_mode_in_response(client: TestClient) -> None:
    xlsx = _make_xlsx({"sites": [["site_code", "country"]]})
    r = _upload(client, xlsx, mode="merge")
    assert r.json()["mode"] == "merge"


# ── POST /export/excel/import — sites ────────────────────────────────────────


def test_import_creates_new_site(client: TestClient) -> None:
    xlsx = _make_xlsx({"sites": [
        ["site_code", "country", "country_code"],
        ["NOBGORiver", "Norway", "NO"],
    ]})
    r = _upload(client, xlsx)
    assert r.status_code == 200
    assert r.json()["summary"]["sites"]["created"] == 1
    assert r.json()["summary"]["sites"]["updated"] == 0

    sites = client.get("/api/sites").json()
    assert any(s["site_code"] == "NOBGORiver" for s in sites)


def test_import_updates_existing_site(client: TestClient) -> None:
    _site(client)
    xlsx = _make_xlsx({"sites": [
        ["site_code", "country", "country_code", "location"],
        ["NOBGORiver", "Norway", "NO", "Updated location"],
    ]})
    r = _upload(client, xlsx)
    assert r.status_code == 200
    assert r.json()["summary"]["sites"]["updated"] == 1

    site = next(s for s in client.get("/api/sites").json() if s["site_code"] == "NOBGORiver")
    assert site["location"] == "Updated location"


def test_import_skips_site_missing_site_code(client: TestClient) -> None:
    xlsx = _make_xlsx({"sites": [
        ["site_code", "country"],
        ["", "Norway"],
    ]})
    r = _upload(client, xlsx, import_valid_rows=True)
    assert r.status_code == 200
    assert r.json()["summary"]["sites"]["skipped"] == 1
    assert r.json()["summary"]["sites"]["created"] == 0


def test_import_skips_site_missing_country(client: TestClient) -> None:
    xlsx = _make_xlsx({"sites": [
        ["site_code", "country"],
        ["NOBGORiver", ""],
    ]})
    r = _upload(client, xlsx, import_valid_rows=True)
    assert r.status_code == 200
    assert r.json()["summary"]["sites"]["skipped"] == 1


# ── POST /export/excel/import — samples ──────────────────────────────────────


def test_import_creates_new_sample(client: TestClient) -> None:
    """A sample whose site_code resolves is created."""
    _site(client)  # creates NOBGORiver, so site_code below resolves
    xlsx = _make_xlsx({"samples": [
        ["sample_code", "site_code", "sample_type", "sampling_date"],
        ["S001", "NOBGORiver", "water", "20240601"],
    ]})
    r = _upload(client, xlsx)
    assert r.status_code == 200
    assert r.json()["summary"]["samples"]["created"] == 1


def test_import_skips_a_sample_whose_site_does_not_resolve(client: TestClient) -> None:
    """Behaviour change (2026-07-30): this row used to be created with site_id NULL.

    sample_code is derived from site + sample_type and is half of the sync merge key
    (UNIQUE (sample_code, sampling_date)), so a sample whose code could not be derived
    correctly produces a key that collides or duplicates across devices. Import now skips the
    row — the same choice seed_samples already made for the same reason — and samples.site_id
    is NOT NULL so the state cannot be reached by any other path either.
    """
    xlsx = _make_xlsx({"samples": [
        ["sample_code", "site_code", "sample_type", "sampling_date"],
        ["S001", "DOES-NOT-EXIST", "water", "20240601"],
    ]})
    r = _upload(client, xlsx, import_valid_rows=True)
    assert r.status_code == 200
    assert r.json()["summary"]["samples"]["created"] == 0
    assert r.json()["summary"]["samples"]["skipped"] == 1


def test_import_skips_a_sample_with_no_sample_type(client: TestClient) -> None:
    """The other half of the derivation: a blank type would give a malformed sample_code."""
    _site(client)
    xlsx = _make_xlsx({"samples": [
        ["sample_code", "site_code", "sample_type", "sampling_date"],
        ["S001", "NOBGORiver", "", "20240601"],
    ]})
    r = _upload(client, xlsx, import_valid_rows=True)
    assert r.status_code == 200
    assert r.json()["summary"]["samples"]["skipped"] == 1


def test_import_skipping_a_sample_does_not_delete_it_in_replace_mode(
    client: TestClient,
) -> None:
    """A skipped row must still count as "seen", or replace mode would prune good data."""
    _site(client)
    created = _upload(client, _make_xlsx({"samples": [
        ["sample_code", "site_code", "sample_type", "sampling_date"],
        ["S001", "NOBGORiver", "water", "20240601"],
    ]}))
    assert created.json()["summary"]["samples"]["created"] == 1

    # Re-import the same key, but with a site_code that no longer resolves.
    r = _upload(client, _make_xlsx({"samples": [
        ["sample_code", "site_code", "sample_type", "sampling_date"],
        ["S001", "DOES-NOT-EXIST", "water", "20240601"],
    ]}), mode="replace", import_valid_rows=True)
    assert r.status_code == 200
    assert r.json()["summary"]["samples"]["skipped"] == 1
    assert len(client.get("/api/samples").json()) == 1


def test_import_updates_existing_sample(client: TestClient) -> None:
    site = _site(client)
    client.post("/api/samples", json={
        "site_id": site["id"], "sampling_date": "20240601", "sample_type": "water",
    })
    # Re-import the same (sample_code, sampling_date) with a changed depth. The code has to be
    # the derived one, and the row needs a resolvable site and a type like any other.
    xlsx = _make_xlsx({"samples": [
        ["sample_code", "site_code", "sample_type", "sampling_date", "depth"],
        ["NOBGORiver_water", "NOBGORiver", "water", "20240601", "5m"],
    ]})
    r = _upload(client, xlsx)
    assert r.status_code == 200, r.text
    assert r.json()["summary"]["samples"]["updated"] == 1
    sample = client.get("/api/samples").json()[0]
    assert sample["depth"] == "5m"


def test_import_skips_sample_missing_sample_code(client: TestClient) -> None:
    xlsx = _make_xlsx({"samples": [
        ["sample_code", "sampling_date"],
        ["", "20240601"],
    ]})
    r = _upload(client, xlsx, import_valid_rows=True)
    assert r.status_code == 200
    assert r.json()["summary"]["samples"]["skipped"] == 1


def test_import_skips_sample_missing_sampling_date(client: TestClient) -> None:
    """A row that carries real data but no sampling_date is a problem, and is reported.

    The row needs a value in a hand-entered column to count as data at all — a row whose only
    populated cell is the auto-filled `sample_code` is a blank template row, covered below.
    """
    xlsx = _make_xlsx({"samples": [
        ["sample_code", "site_code", "sample_type", "sampling_date"],
        ["S001", "NOBGORiver", "water", ""],
    ]})
    r = _upload(client, xlsx, import_valid_rows=True)
    assert r.status_code == 200
    assert r.json()["summary"]["samples"]["skipped"] == 1


def test_import_ignores_rows_holding_only_an_autofilled_value(client: TestClient) -> None:
    """Pre-formatted empty rows are not failed rows.

    In the real curated workbook `sample_code` is a formula column ("Filled out
    automatically - DO NOT EDIT"), and ~530 of the 646 rows on the samples sheet were empty
    rows whose formula still had a cached value. Treating those as data reported 534 problems
    and buried the 6 genuine ones.
    """
    xlsx = _make_xlsx({"samples": [
        ["sample_code", "site_code", "sample_type", "sampling_date"],
        ["S001", "", "", ""],
        ["S002", "", "", ""],
    ]})
    r = _upload(client, xlsx)
    assert r.status_code == 200
    summary = r.json()["summary"]["samples"]
    assert summary["created"] == 0
    assert summary["skipped"] == 0, "blank template rows must not be reported as problems"


# ── POST /export/excel/import — replace mode ──────────────────────────────────


def test_import_replace_mode_deletes_unseen_sites(client: TestClient) -> None:
    _site(client)  # creates NOBGORiver
    xlsx = _make_xlsx({"sites": [
        ["site_code", "country"],
        ["SESTOLake", "Sweden"],  # NOBGORiver not present
    ]})
    r = _upload(client, xlsx, mode="replace")
    assert r.status_code == 200
    assert r.json()["summary"]["sites"]["deleted"] >= 1
    assert r.json()["detail"] == "Import completed (replace mode)"

    site_codes = {s["site_code"] for s in client.get("/api/sites").json()}
    assert "NOBGORiver" not in site_codes
    assert "SESTOLake" in site_codes


def test_import_replace_mode_keeps_site_referenced_by_sample(client: TestClient) -> None:
    site = _site(client)  # creates NOBGORiver
    # Sample references the site — site cannot be deleted while sample exists
    client.post("/api/samples", json={
        "site_id": site["id"], "sampling_date": "20240601", "sample_type": "water",
    })
    # Import a different site; samples sheet not included → existing samples untouched
    xlsx = _make_xlsx({"sites": [
        ["site_code", "country"],
        ["SESTOLake", "Sweden"],
    ]})
    r = _upload(client, xlsx, mode="replace")
    assert r.status_code == 200
    # NOBGORiver protected by its referenced sample
    assert r.json()["summary"]["sites"]["skipped"] >= 1
    site_codes = {s["site_code"] for s in client.get("/api/sites").json()}
    assert "NOBGORiver" in site_codes


def test_import_replace_mode_detail_string(client: TestClient) -> None:
    xlsx = _make_xlsx({"sites": [["site_code", "country"]]})
    r = _upload(client, xlsx, mode="replace")
    assert r.status_code == 200
    assert r.json()["detail"] == "Import completed (replace mode)"


def test_import_merge_mode_detail_string(client: TestClient) -> None:
    xlsx = _make_xlsx({"sites": [["site_code", "country"]]})
    r = _upload(client, xlsx, mode="merge")
    assert r.status_code == 200
    assert r.json()["detail"] == "Import completed"


# ── roundtrip: export → import ────────────────────────────────────────────────


def test_export_then_import_preserves_site(client: TestClient) -> None:
    """A round trip changes nothing, and now says so.

    This used to assert `updated == 1`: re-importing a workbook exported moments earlier
    reported an update to a row where every value was identical. That is not cosmetic —
    `updated_at` drives sync's `incoming_is_newer`, so a round trip made this device's rows
    look newest and able to beat another device's real edits at the next merge.
    """
    _site(client)
    export_bytes = client.get("/api/export/excel").content

    r = _upload(client, export_bytes)
    assert r.status_code == 200
    summary = r.json()["summary"]["sites"]
    assert summary["unchanged"] == 1
    assert summary["updated"] == 0
    assert summary["created"] == 0


def test_reimporting_an_unchanged_row_does_not_restamp_it(
    client: TestClient, db: sqlite3.Connection
) -> None:
    """The consequence the counter stands for: the audit columns must not move."""
    _site(client)
    export_bytes = client.get("/api/export/excel").content
    _upload(client, export_bytes)
    before = db.execute("SELECT updated_at, updated_by FROM sites").fetchone()

    _upload(client, export_bytes)
    after = db.execute("SELECT updated_at, updated_by FROM sites").fetchone()

    assert (after["updated_at"], after["updated_by"]) == (
        before["updated_at"],
        before["updated_by"],
    )


def test_a_genuine_change_in_the_sheet_still_updates_and_restamps(
    client: TestClient, db: sqlite3.Connection
) -> None:
    """Guard against the fix over-reaching."""
    _site(client)
    before = db.execute("SELECT updated_at FROM sites").fetchone()["updated_at"]

    r = _upload(client, _make_xlsx({"sites": [
        ["site_code", "country", "location"],
        ["NOBGORiver", "Norway", "Somewhere new"],
    ]}))
    assert r.status_code == 200, r.text
    assert r.json()["summary"]["sites"]["updated"] == 1
    row = db.execute("SELECT location, updated_at FROM sites").fetchone()
    assert row["location"] == "Somewhere new"
    assert row["updated_at"] != before


def test_import_skips_a_row_with_a_malformed_date_without_failing_the_upload(
    client: TestClient,
) -> None:
    """One bad cell must cost its own row, not the whole import.

    samples.sampling_date carries a YYYYMMDD CHECK. Letting the constraint reject the row
    aborts the transaction, so before this validation a single malformed date in a
    hand-edited sheet returned 409 and rolled back every good row with it.
    """
    _site(client)
    xlsx = _make_xlsx({"samples": [
        ["sample_code", "site_code", "sample_type", "sampling_date"],
        ["S001", "NOBGORiver", "water", "20240601"],
        ["S002", "NOBGORiver", "water", "1 June 2024"],
    ]})
    r = _upload(client, xlsx, import_valid_rows=True)
    assert r.status_code == 200, r.text
    assert r.json()["summary"]["samples"]["created"] == 1
    assert r.json()["summary"]["samples"]["skipped"] == 1
    codes = {s["sample_code"] for s in client.get("/api/samples").json()}
    assert codes == {"S001"}


def test_import_skips_a_row_with_a_malformed_date_extraction(client: TestClient) -> None:
    _site(client)
    xlsx = _make_xlsx({"samples": [
        ["sample_code", "site_code", "sample_type", "sampling_date", "date_extraction"],
        ["S001", "NOBGORiver", "water", "20240601", "2024-06-01"],
    ]})
    r = _upload(client, xlsx, import_valid_rows=True)
    assert r.status_code == 200, r.text
    assert r.json()["summary"]["samples"]["skipped"] == 1


# ── POST /export/excel/import — reading real-world workbooks ──────────────────
#
# Each case here cost an entire sheet before: the curated metadata workbooks put the header
# on row 2 behind an annotation row, key the site by site_ID rather than site_code, vary the
# case of column names, and carry a prose description row under the header.


def test_import_accepts_site_id_as_a_spelling_of_site_code(client: TestClient) -> None:
    """The seed CSV format and the hand-maintained workbooks both use site_ID."""
    xlsx = _make_xlsx({"sites": [
        ["site_ID", "country", "country_code"],
        ["NOBGORiver", "Norway", "NO"],
    ]})
    r = _upload(client, xlsx)
    assert r.status_code == 200, r.text
    assert r.json()["summary"]["sites"]["created"] == 1
    assert any(s["site_code"] == "NOBGORiver" for s in client.get("/api/sites").json())


def test_import_matches_column_names_case_insensitively(client: TestClient) -> None:
    """`Biomeme_sample_ID` is the real workbook's spelling of biomeme_sample_id."""
    _site(client)
    client.post("/api/samples", json={
        "site_id": client.get("/api/sites").json()[0]["id"],
        "sample_type": "water", "sampling_date": "20240601",
    })
    xlsx = _make_xlsx({"biomeme": [
        ["biomeme_run_name", "Sample_Code", "Sampling_Date", "Biomeme_sample_ID"],
        ["BM-001", "NOBGORiver_water", "20240601", "BS-9"],
    ]})
    r = _upload(client, xlsx)
    assert r.status_code == 200, r.text
    runs = client.get("/api/biomeme-runs").json()
    assert runs and runs[0]["biomeme_sample_id"] == "BS-9"


def test_import_finds_a_header_below_an_annotation_row(client: TestClient) -> None:
    """nanopore/biomeme sheets in the curated template have the header on row 2."""
    xlsx = _make_xlsx({"sites": [
        ["from 'somewhere' sheet", None, None],
        ["site_code", "country", "country_code"],
        ["NOBGORiver", "Norway", "NO"],
    ]})
    r = _upload(client, xlsx)
    assert r.status_code == 200, r.text
    assert r.json()["summary"]["sites"]["created"] == 1


def test_import_drops_the_template_description_row(client: TestClient) -> None:
    """The row of prose under the header is template furniture, not a failed row."""
    _site(client)
    xlsx = _make_xlsx({"samples": [
        ["sample_code", "site_code", "sample_type", "sampling_date"],
        ["Filled out automatically - DO NOT EDIT", "Select from drop-down",
         "What type of sample", "Enter the date (yyyymmdd)"],
        ["S001", "NOBGORiver", "water", "20240601"],
    ]})
    r = _upload(client, xlsx)
    assert r.status_code == 200, r.text
    assert r.json()["summary"]["samples"]["created"] == 1
    # The description row must not be counted as a skipped data row.
    assert r.json()["summary"]["samples"]["skipped"] == 0


def test_import_still_reports_a_later_malformed_row(client: TestClient) -> None:
    """Only the row directly under the header gets description treatment."""
    _site(client)
    xlsx = _make_xlsx({"samples": [
        ["sample_code", "site_code", "sample_type", "sampling_date"],
        ["S001", "NOBGORiver", "water", "20240601"],
        ["S002", "NOBGORiver", "water", "not-a-date"],
    ]})
    r = _upload(client, xlsx, import_valid_rows=True)
    assert r.status_code == 200, r.text
    assert r.json()["summary"]["samples"]["created"] == 1
    assert r.json()["summary"]["samples"]["skipped"] == 1


def test_import_ignores_a_sheet_with_no_recognisable_header(client: TestClient) -> None:
    xlsx = _make_xlsx({"sites": [
        ["totally", "unrelated", "columns"],
        ["a", "b", "c"],
    ]})
    r = _upload(client, xlsx)
    assert r.status_code == 200, r.text
    assert r.json()["summary"]["sites"]["created"] == 0


# ── POST /export/excel/import — the rejection contract ────────────────────────
#
# The default is all-or-nothing with an explanation: an upload containing any unimportable row
# is rejected, nothing is written, and the response names every offending row and why. Partial
# import happens only when the caller opts in, which the UI will do once the operator has seen
# the list. Before this, the only signal was a `skipped` integer — and for a real curated
# workbook the response read "0 created, 0 skipped" while importing nothing at all.


def test_import_rejects_the_whole_upload_when_a_row_is_unimportable(
    client: TestClient,
) -> None:
    _site(client)
    xlsx = _make_xlsx({"samples": [
        ["sample_code", "site_code", "sample_type", "sampling_date"],
        ["S001", "NOBGORiver", "water", "20240601"],
        ["S002", "NOSUCHSITE", "water", "20240601"],
    ]})
    r = _upload(client, xlsx)
    assert r.status_code == 422, r.text
    # Nothing was written, not even the valid row.
    assert client.get("/api/samples").json() == []


def test_the_rejection_names_the_sheet_row_and_reason(client: TestClient) -> None:
    _site(client)
    xlsx = _make_xlsx({"samples": [
        ["sample_code", "site_code", "sample_type", "sampling_date"],
        ["S001", "NOBGORiver", "water", "20240601"],
        ["S002", "NOSUCHSITE", "water", "20240601"],
        ["S003", "NOBGORiver", "water", "nonsense"],
    ]})
    detail = _upload(client, xlsx).json()["detail"]

    assert detail["problem_count"] == 2
    problems = {p["row"]: p for p in detail["problems"]}
    assert set(problems) == {3, 4}, problems
    assert problems[3]["sheet"] == "samples"
    assert "NOSUCHSITE" in problems[3]["reason"]
    assert "nonsense" in problems[4]["reason"]
    # The summary of what *would* have happened travels with the rejection.
    assert detail["summary"]["samples"]["created"] == 1


def test_opting_in_imports_the_valid_rows_and_still_lists_the_problems(
    client: TestClient,
) -> None:
    _site(client)
    xlsx = _make_xlsx({"samples": [
        ["sample_code", "site_code", "sample_type", "sampling_date"],
        ["S001", "NOBGORiver", "water", "20240601"],
        ["S002", "NOSUCHSITE", "water", "20240601"],
    ]})
    r = _upload(client, xlsx, import_valid_rows=True)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["summary"]["samples"]["created"] == 1
    assert body["problem_count"] == 1
    assert "NOSUCHSITE" in body["problems"][0]["reason"]
    assert len(client.get("/api/samples").json()) == 1


def test_a_clean_upload_reports_no_problems(client: TestClient) -> None:
    _site(client)
    xlsx = _make_xlsx({"samples": [
        ["sample_code", "site_code", "sample_type", "sampling_date"],
        ["S001", "NOBGORiver", "water", "20240601"],
    ]})
    body = _upload(client, xlsx).json()
    assert body["problem_count"] == 0
    assert body["problems"] == []


def test_an_unknown_sample_type_is_reported_like_any_unresolved_reference(
    client: TestClient,
) -> None:
    """lookup_values references are the same class of problem as site_code references.

    They cannot be foreign keys (lookup_values is keyed (list, code)), and the bulk paths must
    not be a way around the 422 the API returns for an unknown code.
    """
    _site(client)
    xlsx = _make_xlsx({"samples": [
        ["sample_code", "site_code", "sample_type", "sampling_date"],
        ["S001", "NOBGORiver", "not-a-real-type", "20240601"],
    ]})
    detail = _upload(client, xlsx).json()["detail"]
    assert detail["problem_count"] == 1
    assert "not-a-real-type" in detail["problems"][0]["reason"]
    assert "lookup" in detail["problems"][0]["reason"]


# ── POST /export/excel/import?dry_run=true — preview ─────────────────────────
#
# A preview is the real import rolled back, so what it reports is exactly what applying
# would do. These tests hold that equivalence, because the value of the preview rests on it.


def _preview(client: TestClient, xlsx: bytes, mode: str = "merge") -> dict:
    r = client.post(
        "/api/export/excel/import",
        files={"file": ("test.xlsx", xlsx, _XLSX_CONTENT_TYPE)},
        params={"mode": mode, "dry_run": True},
    )
    assert r.status_code == 200, r.text
    return r.json()


def test_preview_writes_nothing(client: TestClient) -> None:
    _site(client)
    body = _preview(client, _make_xlsx({"samples": [
        ["sample_code", "site_code", "sample_type", "sampling_date"],
        ["S001", "NOBGORiver", "water", "20240601"],
    ]}))
    assert body["dry_run"] is True
    assert body["summary"]["samples"]["created"] == 1
    assert client.get("/api/samples").json() == [], "a preview must not write"


def test_preview_is_never_an_error_however_bad_the_workbook(client: TestClient) -> None:
    """Unlike an apply, a preview answers the question rather than refusing it."""
    body = _preview(client, _make_xlsx({"samples": [
        ["sample_code", "site_code", "sample_type", "sampling_date"],
        ["S001", "NOSUCHSITE", "water", "nonsense"],
    ]}))
    assert body["problem_count"] == 1
    assert body["summary"]["samples"]["skipped"] == 1


def test_preview_reports_what_applying_actually_does(client: TestClient) -> None:
    """The equivalence the whole design rests on: same numbers, same problems."""
    _site(client)
    xlsx = _make_xlsx({"samples": [
        ["sample_code", "site_code", "sample_type", "sampling_date"],
        ["S001", "NOBGORiver", "water", "20240601"],
        ["S002", "NOSUCHSITE", "water", "20240601"],
    ]})

    previewed = _preview(client, xlsx)
    applied = _upload(client, xlsx, import_valid_rows=True).json()

    assert previewed["summary"] == applied["summary"]
    assert previewed["problems"] == applied["problems"]


def test_preview_distinguishes_unchanged_from_updated(client: TestClient) -> None:
    """What the operator most needs to know: how much of this would actually change."""
    _site(client)
    xlsx = _make_xlsx({"sites": [
        ["site_code", "country", "location"],
        ["NOBGORiver", "Norway", "Somewhere new"],
    ]})
    _upload(client, xlsx)  # apply once

    body = _preview(client, xlsx)  # the same workbook again
    assert body["summary"]["sites"]["unchanged"] == 1
    assert body["summary"]["sites"]["updated"] == 0


def test_preview_respects_replace_mode(client: TestClient) -> None:
    """Replace mode deletes; the preview has to say so before the operator commits to it."""
    _site(client)
    body = _preview(client, _make_xlsx({"sites": [
        ["site_code", "country"],
        ["SESTOLake", "Sweden"],
    ]}), mode="replace")
    assert body["summary"]["sites"]["deleted"] >= 1
    assert len(client.get("/api/sites").json()) == 1, "preview must not delete either"


def test_a_first_data_row_with_one_bad_value_is_reported_not_mistaken_for_furniture(
    client: TestClient,
) -> None:
    """The boundary of the description-row heuristic.

    A template description row is prose in every cell. A data row with one malformed value is
    a problem to report. The first version of this heuristic could not tell them apart — it
    fired on any first row whose sampling_date was not 8 digits, so that row disappeared with
    no report.
    """
    _site(client)
    r = _upload(client, _make_xlsx({"samples": [
        ["sample_code", "site_code", "sample_type", "sampling_date"],
        ["S001", "NOBGORiver", "water", "nonsense"],
    ]}), import_valid_rows=True)

    assert r.status_code == 200, r.text
    assert r.json()["summary"]["samples"]["skipped"] == 1
    assert "nonsense" in r.json()["problems"][0]["reason"]


# ── POST /export/excel/import — the API rules the importer used to bypass ─────
# The endpoints enforce two site rules the importer did not: the shared code
# vocabulary (_check_code_conflicts — 'NO' must not mean Norway in one row and
# Norfolk in another) and the derived-code freeze (country_code/city_code/site
# are immutable once samples reference the site, because every sample's stored
# sample_code — half the sync merge key, and the name of pipeline output on
# disk — is derived from them). A workbook must obey the same rules.


def test_import_does_not_bind_a_country_code_to_a_second_name(client: TestClient) -> None:
    _site(client)  # binds NO = Norway
    xlsx = _make_xlsx({"sites": [
        ["site_code", "country", "country_code"],
        ["ZZNEW01", "Norfolk", "NO"],
    ]})
    r = _upload(client, xlsx, import_valid_rows=True)

    assert r.status_code == 200, r.text
    assert r.json()["summary"]["sites"]["created"] == 0
    assert r.json()["summary"]["sites"]["skipped"] == 1
    assert "already associated" in r.json()["problems"][0]["reason"]
    assert not any(s["site_code"] == "ZZNEW01" for s in client.get("/api/sites").json())


def test_import_conflict_check_guards_updates_too(client: TestClient) -> None:
    _site(client)  # binds NO = Norway
    other = _site(client, country="Sweden", country_code="SE", city_code="GOT", site="Dock")
    xlsx = _make_xlsx({"sites": [
        ["site_code", "country", "country_code"],
        [other["site_code"], "Norfolk", "NO"],
    ]})
    r = _upload(client, xlsx, import_valid_rows=True)

    assert r.json()["summary"]["sites"]["skipped"] == 1
    assert "already associated" in r.json()["problems"][0]["reason"]
    unchanged = next(s for s in client.get("/api/sites").json() if s["id"] == other["id"])
    assert unchanged["country"] == "Sweden"


def test_import_cannot_change_frozen_components_of_a_site_with_samples(
    client: TestClient,
) -> None:
    site = _site(client)
    client.post("/api/samples", json={
        "site_id": site["id"], "sampling_date": "20240601", "sample_type": "water",
    })
    xlsx = _make_xlsx({"sites": [
        ["site_code", "country", "city_code"],
        ["NOBGORiver", "Norway", "OSL"],
    ]})
    r = _upload(client, xlsx, import_valid_rows=True)

    assert r.json()["summary"]["sites"]["skipped"] == 1
    assert "sample" in r.json()["problems"][0]["reason"]
    unchanged = next(s for s in client.get("/api/sites").json() if s["id"] == site["id"])
    assert unchanged["city_code"] == "BGO"


def test_import_can_change_components_of_a_site_with_no_samples(client: TestClient) -> None:
    # Mirrors the endpoint rule: with no stored copies anywhere, correcting a typo is allowed.
    site = _site(client)
    xlsx = _make_xlsx({"sites": [
        ["site_code", "country", "city_code"],
        ["NOBGORiver", "Norway", "OSL"],
    ]})
    r = _upload(client, xlsx, import_valid_rows=True)

    assert r.json()["summary"]["sites"]["updated"] == 1
    updated = next(s for s in client.get("/api/sites").json() if s["id"] == site["id"])
    assert updated["city_code"] == "OSL"


def test_import_leaves_site_columns_absent_from_the_sheet_untouched(
    client: TestClient,
) -> None:
    # A sheet that never mentions a column says nothing about it. Treating absence as
    # "clear" would let the real curated workbook — sites keyed by site_ID + country
    # only — wipe the components of every matched site (and trip the freeze above for
    # any site with samples). A present-but-blank cell still clears, as everywhere else.
    site = _site(client, location="Harbour")
    xlsx = _make_xlsx({"sites": [
        ["site_code", "country", "location"],
        ["NOBGORiver", "Norway", "Updated location"],
    ]})
    r = _upload(client, xlsx, import_valid_rows=True)

    assert r.json()["summary"]["sites"]["updated"] == 1
    updated = next(s for s in client.get("/api/sites").json() if s["id"] == site["id"])
    assert updated["location"] == "Updated location"
    assert updated["city_code"] == "BGO"
    assert updated["site"] == "River"


# ── Absent columns leave fields alone on every sheet ──────────────────────────
# The same defect fixed for sites existed on the other three sheets: update
# candidates built with r.get() read an absent column as None and cleared the
# stored value. A sheet that never mentions a column says nothing about it.


def test_import_leaves_sample_columns_absent_from_the_sheet_untouched(
    client: TestClient,
) -> None:
    site = _site(client)
    created = client.post("/api/samples", json={
        "site_id": site["id"], "sampling_date": "20240601", "sample_type": "water",
        "depth": "5m", "comments": "keep me",
    }).json()
    xlsx = _make_xlsx({"samples": [
        ["sample_code", "site_code", "sample_type", "sampling_date"],
        [created["sample_code"], "NOBGORiver", "water", "20240601"],
    ]})
    r = _upload(client, xlsx, import_valid_rows=True)

    assert r.status_code == 200, r.text
    after = next(s for s in client.get("/api/samples").json() if s["id"] == created["id"])
    assert after["depth"] == "5m"
    assert after["comments"] == "keep me"


def test_import_leaves_nanopore_columns_absent_from_the_sheet_untouched(
    client: TestClient,
) -> None:
    site = _site(client)
    sample = client.post("/api/samples", json={
        "site_id": site["id"], "sampling_date": "20240601", "sample_type": "water",
    }).json()
    created = client.post("/api/nanopore-runs", json={
        "run_accession": "ERR900", "barcode": "barcode01", "sample_id": sample["id"],
        "protocol_id": "SQK-LSK114", "comments": "run-level note",
    }).json()
    assert created["sample_id"] == sample["id"]

    xlsx = _make_xlsx({"nanopore": [
        ["run_accession", "barcode"],
        ["ERR900", "barcode01"],
    ]})
    r = _upload(client, xlsx, import_valid_rows=True)

    assert r.status_code == 200, r.text
    after = next(n for n in client.get("/api/nanopore-runs").json() if n["id"] == created["id"])
    assert after["sample_id"] == sample["id"]      # the link survives
    assert after["protocol_id"] == "SQK-LSK114"    # accession fields survive
    assert after["comments"] == "run-level note"


def test_import_leaves_biomeme_columns_absent_from_the_sheet_untouched(
    client: TestClient,
) -> None:
    created = client.post("/api/biomeme-runs", json={
        "biomeme_run_name": "BB_20240210_run1", "comments": "keep me",
    }).json()
    # Two known columns, because a sheet needs that many for its header to be recognised —
    # a single-column sheet is ignored outright, which made the first version of this test
    # pass without importing anything.
    xlsx = _make_xlsx({"biomeme": [
        ["biomeme_run_name", "dilution_factor"],
        ["BB_20240210_run1", 10],
    ]})
    r = _upload(client, xlsx, import_valid_rows=True)

    assert r.status_code == 200, r.text
    assert r.json()["summary"]["biomeme"]["updated"] == 1
    after = next(b for b in client.get("/api/biomeme-runs").json() if b["id"] == created["id"])
    assert after["dilution_factor"] == 10
    assert after["comments"] == "keep me"


def test_replace_mode_keeps_a_site_the_freeze_refused_to_update(client: TestClient) -> None:
    # A row the freeze skips is still marked "seen", so replace mode must not prune the
    # site it refused to touch — a skip protects data, never deletes it.
    site = _site(client)
    client.post("/api/samples", json={
        "site_id": site["id"], "sampling_date": "20240601", "sample_type": "water",
    })
    xlsx = _make_xlsx({"sites": [
        ["site_code", "country", "city_code"],
        ["NOBGORiver", "Norway", "OSL"],
    ]})
    r = _upload(client, xlsx, mode="replace", import_valid_rows=True)

    assert r.status_code == 200, r.text
    assert r.json()["summary"]["sites"]["skipped"] == 1
    assert r.json()["summary"]["sites"]["deleted"] == 0
    survivor = next(s for s in client.get("/api/sites").json() if s["id"] == site["id"])
    assert survivor["city_code"] == "BGO"


# ── Preview lists WHAT changes, not only how many ─────────────────────────────
# The importer computed the changed fields per row (that is how "identical" is
# counted) and threw them away; the preview dialog could only show counts.


def test_preview_lists_the_changed_fields_per_row(client: TestClient) -> None:
    _site(client, location="Old location")
    xlsx = _make_xlsx({"sites": [
        ["site_code", "country", "location"],
        ["NOBGORiver", "Norway", "New location"],
    ]})
    r = client.post(
        "/api/export/excel/import",
        files={"file": ("t.xlsx", xlsx, _XLSX_CONTENT_TYPE)},
        params={"mode": "merge", "dry_run": True},
    )

    assert r.status_code == 200, r.text
    body = r.json()
    assert body["change_count"] == 1
    change = body["changes"][0]
    assert change["sheet"] == "sites"
    assert change["row"] == 2
    assert change["fields"] == [
        {"field": "location", "old": "Old location", "new": "New location"}
    ]


def test_preview_reports_no_changes_for_an_identical_row(client: TestClient) -> None:
    _site(client)
    xlsx = _make_xlsx({"sites": [
        ["site_code", "country", "country_code"],
        ["NOBGORiver", "Norway", "NO"],
    ]})
    r = client.post(
        "/api/export/excel/import",
        files={"file": ("t.xlsx", xlsx, _XLSX_CONTENT_TYPE)},
        params={"mode": "merge", "dry_run": True},
    )

    assert r.json()["change_count"] == 0
    assert r.json()["changes"] == []


def test_preview_records_run_level_nanopore_changes(client: TestClient) -> None:
    # The accession update does not go through apply_update, so it needs its own
    # recording — and it is the change the summary cannot even count (an
    # accession-only edit leaves the barcode row "identical").
    client.post("/api/nanopore-runs", json={
        "run_accession": "ERR900", "barcode": "barcode01", "comments": "old note",
    })
    xlsx = _make_xlsx({"nanopore": [
        ["run_accession", "barcode", "comments"],
        ["ERR900", "barcode01", "new note"],
    ]})
    r = client.post(
        "/api/export/excel/import",
        files={"file": ("t.xlsx", xlsx, _XLSX_CONTENT_TYPE)},
        params={"mode": "merge", "dry_run": True},
    )

    body = r.json()
    assert body["change_count"] == 1
    assert body["changes"][0]["fields"] == [
        {"field": "comments", "old": "old note", "new": "new note"}
    ]
