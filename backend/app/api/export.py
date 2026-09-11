import io
import sqlite3
import uuid
from dataclasses import dataclass, field
from typing import Any, Literal, Optional

import openpyxl
from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl.styles import Font

from ..database import get_db
from ..utils import build_update, is_yyyymmdd, resolve_sample_id
from ..utils import utc_now_str as _now
from .sites import find_code_conflict

router = APIRouter(prefix="/export", tags=["export"])


def _bold_header(ws) -> None:
    for cell in ws[1]:
        cell.font = Font(bold=True)


# At most this many problems are returned. A genuinely broken file could otherwise produce
# thousands, and no UI can show them usefully. The total is always reported separately so the
# cap never hides how much was wrong.
_MAX_REPORTED_PROBLEMS = 200

# Which table each sheet upserts into, for the shared difference-aware update.
_SHEET_TABLES = {
    "sites": "sites",
    "samples": "samples",
    "nanopore": "nanopore_runs",
    "biomeme": "biomeme_runs",
}


@dataclass
class _ImportReport:
    """Counters plus the reason for every row that was not imported.

    Skipping and explaining are one operation here on purpose. Previously a skip only bumped
    an integer, so the operator was told "12 skipped" with no way to learn which rows or why
    — and for a real workbook it read "0 created, 0 skipped" while importing nothing at all.
    """

    summary: dict
    problems: list[dict] = field(default_factory=list)
    changes: list[dict] = field(default_factory=list)

    def skip(self, sheet: str, row: Optional[int], reason: str) -> None:
        self.summary[sheet]["skipped"] += 1
        self.problems.append({"sheet": sheet, "row": row, "reason": reason})

    def record_change(
        self, sheet: str, row: Optional[int], existing: sqlite3.Row, changed: dict
    ) -> None:
        """Remember WHICH fields a row changes, so the preview can show more than counts.

        Merged per (sheet, row): the nanopore sheet writes a row's barcode-level and
        run-level fields separately, and the operator thinks of them as one row.
        """
        fields = [
            {"field": key, "old": existing[key], "new": value}
            for key, value in changed.items()
        ]
        for entry in self.changes:
            if entry["sheet"] == sheet and entry["row"] == row:
                entry["fields"].extend(fields)
                return
        self.changes.append({"sheet": sheet, "row": row, "fields": fields})

    def apply_update(
        self,
        db: sqlite3.Connection,
        sheet: str,
        existing: sqlite3.Row,
        candidates: dict,
        now: str,
        actor: str,
        row: Optional[int] = None,
    ) -> None:
        """Write only the columns the sheet actually changes, and count honestly.

        The import used to issue an unconditional full-row UPDATE whenever the key matched, so
        re-importing a byte-identical workbook reported "updated: N" and moved `updated_at` on
        every matched row. That is the same presence-vs-difference defect fixed in the update
        endpoints, and it is worse here because of what follows: `updated_at` drives sync's
        `incoming_is_newer`, so a no-op import made this device's copies look newest and could
        beat another device's genuine edits at the next merge.

        Rows that differ in nothing are now counted as `unchanged` rather than `updated`, which
        also gives the planned preview something truthful to show.
        """
        changed = {key: value for key, value in candidates.items() if existing[key] != value}
        if not changed:
            self.summary[sheet]["unchanged"] += 1
            return
        sql, params = build_update(_SHEET_TABLES[sheet], changed, existing["id"], now, actor)
        db.execute(sql, params)
        self.summary[sheet]["updated"] += 1
        self.record_change(sheet, row, existing, changed)

    def as_detail(self, message: str) -> dict:
        return {
            "message": message,
            "problem_count": len(self.problems),
            "problems": self.problems[:_MAX_REPORTED_PROBLEMS],
            "problems_truncated": max(0, len(self.problems) - _MAX_REPORTED_PROBLEMS),
            "summary": self.summary,
        }


# Every column name the importer understands, in the spelling the import functions use.
_CANONICAL_FIELDS = (
    # sites
    "site_code", "site", "country", "country_code", "city_code", "city",
    "location", "longitude", "latitude", "comments",
    # samples
    "sample_code", "sample_type", "sampling_date", "depth", "elevation",
    "partner_sample_code", "nucleic_acid_concentration", "extract_volume",
    "elution_volume", "date_extraction", "comments_sampling", "comments_extraction",
    # nanopore
    "run_accession", "barcode", "protocol_id", "sequencing_kit_id", "type",
    "runName", "sampleName",
    # biomeme
    "biomeme_run_name", "biomeme_sample_id", "dilution_factor",
)

# Alternative spellings accepted on import. `site_ID` is what the seed CSV format and the
# hand-maintained metadata workbooks use, while the Excel *export* emits `site_code` — the two
# ingestion paths disagreed, so a real curated workbook resolved no sites at all. Export is
# deliberately unchanged: it keeps emitting the canonical names.
_COLUMN_ALIASES = {"site_id": "site_code"}

# lowercased header -> canonical field name. Matching is case-insensitive because the
# workbooks also vary case (e.g. `Biomeme_sample_ID` for `biomeme_sample_id`), which silently
# dropped those columns before.
_HEADER_LOOKUP: dict[str, str] = {f.lower(): f for f in _CANONICAL_FIELDS}
_HEADER_LOOKUP.update(_COLUMN_ALIASES)

# How many rows to search for the header. The curated workbooks put an annotation row above
# it ("from 'samples' sheet", "library_prep_details"), so it is not always row 1.
_HEADER_SEARCH_DEPTH = 10

# A row is the header if at least this many of its cells name a field we understand. Two
# avoids mistaking a single stray label in an annotation row for the header.
_MIN_HEADER_MATCHES = 2


def _canonical_header(value: Any) -> str:
    """Map one header cell to a canonical field name, or "" if it names nothing we know."""
    if value is None:
        return ""
    text = str(value).strip()
    return _HEADER_LOOKUP.get(text.lower(), text)


def _find_header_row(rows: list[tuple]) -> int:
    """Index of the header row within *rows*, or -1 when no row looks like one."""
    for index, row in enumerate(rows[:_HEADER_SEARCH_DEPTH]):
        known = sum(
            1 for cell in row if _canonical_header(cell) in _HEADER_LOOKUP.values()
        )
        if known >= _MIN_HEADER_MATCHES:
            return index
    return -1


# Columns that the workbook fills in by formula rather than by hand, per sheet. A row whose
# only populated cells are these carries no information: it is a pre-formatted empty row whose
# auto-fill formula still has a cached value. The real curated workbook has ~530 of them on the
# samples sheet, and counting them as data would bury the handful of genuine problems.
#
# Deliberately per-sheet: `sample_code` is auto-filled on the samples sheet but is the
# identifying reference on nanopore/biomeme, and `site_code` is the identifying key on sites.
_AUTOFILLED_COLUMNS: dict[str, frozenset[str]] = {
    "sites": frozenset(),
    "samples": frozenset({"sample_code"}),
    "nanopore": frozenset({"alias", "minknow_sample_id", "sample_id"}),
    "biomeme": frozenset(),
}


def _sheet_rows_by_header(ws, autofilled: frozenset[str] = frozenset()) -> list[dict[str, Any]]:
    """Return the sheet's data rows as canonical-keyed dicts, with 1-based row numbers.

    Three things this has to cope with in real workbooks, each of which silently cost the
    whole sheet before:

    * **The header is not always row 1.** The nanopore and biomeme sheets of the curated
      template carry an annotation row above it, so the header is located rather than assumed.
    * **A human-readable description row often follows the header** ("Enter the date of
      sample collection (yyyymmdd)"). The seed CSV reader already drops it; only the row
      *immediately* after the header is treated this way, so a later malformed row is still
      reported as a problem instead of being quietly ignored.
    * **Declared width, not used width.** These sheets declare 16 384 columns, so iterating
      every cell cost 74 s for one sheet. Iteration is bounded to the header's width.

    ``__row__`` carries the spreadsheet row number so problems can be reported by row.
    """
    # Read only enough rows to locate the header. Full width here is cheap because it is a
    # handful of rows; reading the *whole* sheet at full width is what cost 74 s.
    head = list(ws.iter_rows(min_row=1, max_row=_HEADER_SEARCH_DEPTH, values_only=True))
    if not head:
        return []

    header_index = _find_header_row(head)
    if header_index < 0:
        return []

    header = [_canonical_header(c) for c in head[header_index]]
    # Trailing unnamed columns are dropped, and the resulting width then bounds the scan of
    # the body — so a sheet declaring 16 384 columns is read at its real width instead.
    width = max((i + 1 for i, h in enumerate(header) if h), default=0)
    header = header[:width]
    if not width:
        return []

    body = ws.iter_rows(min_row=header_index + 2, max_col=width, values_only=True)

    informative = [h for h in header if h and h not in autofilled]

    out: list[dict[str, Any]] = []
    for offset, row in enumerate(body):
        values = list(row)
        rec: dict[str, Any] = {}
        for i, h in enumerate(header):
            if not h:
                continue
            rec[h] = values[i] if i < len(values) else None
        # Blank when nothing outside the auto-filled columns is populated.
        if not any(_str(rec.get(h)) is not None for h in informative):
            continue
        rec["__row__"] = header_index + 2 + offset
        # The description row sits directly beneath the header and never anywhere else.
        if offset == 0 and _looks_like_description_row(rec):
            continue
        out.append(rec)
    return out


# A cell is prose when it reads like an instruction rather than a value: long enough to be a
# sentence, and containing a space. "Enter the date of sample collection (yyyymmdd)" qualifies;
# "NOSUCHSITE", "nonsense" and "20240601" do not.
_PROSE_MIN_CHARS = 25

# How many prose cells make a row template furniture. Two, because one malformed value is a
# data error to report — not a reason to discard the row.
_MIN_PROSE_CELLS = 2


def _looks_like_description_row(rec: dict[str, Any]) -> bool:
    """True when a row is template prose rather than data.

    The templates put a line of instructions under the header, and every one of its cells is a
    sentence. Detection counts those rather than testing a single field's format, which was the
    first attempt: it treated any first data row whose sampling_date was malformed as
    furniture, so that row vanished with no report — the silent behaviour this whole area is
    meant to be rid of. A test caught it.

    Only ever applied to the row directly below the header.
    """
    prose_cells = 0
    for key, value in rec.items():
        if key == "__row__":
            continue
        text = _str(value)
        if text is not None and len(text) >= _PROSE_MIN_CHARS and " " in text:
            prose_cells += 1
    return prose_cells >= _MIN_PROSE_CELLS


def _str(v: Any) -> Optional[str]:
    if v is None:
        return None
    s = str(v).strip()
    return s if s else None


def _flt(v: Any) -> Optional[float]:
    s = _str(v)
    if s is None:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def _lookup_code_exists(db: sqlite3.Connection, list_name: str, code: Optional[str]) -> bool:
    """True when *code* exists in ``lookup_values`` for *list_name* (or is absent).

    These columns — sample_type, protocol_id, sequencing_kit_id, mpox type — are references
    exactly like site_id is a reference; they simply cannot be declared as foreign keys because
    ``lookup_values`` is keyed ``(list, code)``. The API rejects an unknown code with 422, so the
    bulk paths must not be a way around that.
    """
    if code is None:
        return True
    return (
        db.execute(
            "SELECT 1 FROM lookup_values WHERE list = ? AND code = ?", (list_name, code)
        ).fetchone()
        is not None
    )


def _first_unknown_lookup(
    db: sqlite3.Connection, pairs: tuple[tuple[str, Optional[str]], ...]
) -> Optional[tuple[str, str]]:
    """Return the first ``(list_name, code)`` that does not resolve, or None."""
    for list_name, code in pairs:
        if not _lookup_code_exists(db, list_name, code):
            return list_name, code or ""
    return None


def _find_site_id(db: sqlite3.Connection, site_code: Optional[str]) -> Optional[str]:
    if not site_code:
        return None
    row = db.execute("SELECT id FROM sites WHERE site_code = ?", (site_code,)).fetchone()
    return row["id"] if row else None


@router.get("/excel", summary="Export metadata to Excel")
def export_excel(db: sqlite3.Connection = Depends(get_db)):
    sites = db.execute("SELECT * FROM sites").fetchall()
    samples = db.execute("SELECT * FROM samples").fetchall()
    nanopore_runs = db.execute(
        "SELECT nr.*, nra.run_accession, s.sample_code, s.sampling_date, "
        "nra.protocol_id, nra.sequencing_kit_id, nr.type, nra.runName, nra.sampleName, nra.comments, "
        "CASE WHEN s.sample_code IS NOT NULL"
        "          AND nra.protocol_id IS NOT NULL"
        "          AND nra.sequencing_kit_id IS NOT NULL"
        "     THEN s.sample_code || '_' || nra.protocol_id || '_' || nra.sequencing_kit_id"
        "     ELSE NULL END AS minknow_sample_id, "
        "CASE WHEN s.sample_code IS NOT NULL"
        "     THEN s.sample_code || '_' || nr.barcode"
        "     ELSE NULL END AS alias "
        "FROM nanopore_runs nr "
        "LEFT JOIN samples s ON s.id = nr.sample_id "
        "LEFT JOIN nanopore_run_accessions nra ON nra.id = nr.accession_id"
    ).fetchall()
    biomeme_runs = db.execute(
        "SELECT br.*, s.sample_code, s.sampling_date FROM biomeme_runs br "
        "LEFT JOIN samples s ON s.id = br.sample_id"
    ).fetchall()

    sites_by_id = {r["id"]: r["site_code"] for r in sites}

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("sites")
    ws.append(
        [
            "site_code",
            "site",
            "country",
            "country_code",
            "city_code",
            "city",
            "location",
            "longitude",
            "latitude",
            "comments",
        ]
    )
    _bold_header(ws)
    for s in sites:
        ws.append(
            [
                s["site_code"],
                s["site"],
                s["country"],
                s["country_code"],
                s["city_code"],
                s["city"],
                s["location"],
                s["longitude"],
                s["latitude"],
                s["comments"],
            ]
        )

    ws = wb.create_sheet("samples")
    ws.append(
        [
            "sample_code",
            "site_code",
            "sample_type",
            "sampling_date",
            "depth",
            "elevation",
            "partner_sample_code",
            "nucleic_acid_concentration",
            "extract_volume",
            "elution_volume",
            "date_extraction",
            "comments_sampling",
            "comments_extraction",
            "comments",
        ]
    )
    _bold_header(ws)
    for s in samples:
        ws.append(
            [
                s["sample_code"],
                sites_by_id.get(s["site_id"], ""),
                s["sample_type"],
                s["sampling_date"],
                s["depth"],
                s["elevation"],
                s["partner_sample_code"],
                s["nucleic_acid_concentration"],
                s["extract_volume"],
                s["elution_volume"],
                s["date_extraction"],
                s["comments_sampling"],
                s["comments_extraction"],
                s["comments"],
            ]
        )

    ws = wb.create_sheet("nanopore")
    ws.append(
        [
            "run_accession",
            "sample_code",
            "sampling_date",
            "barcode",
            "minknow_sample_id",
            "alias",
            "protocol_id",
            "sequencing_kit_id",
            "type",
            "runName",
            "sampleName",
            "comments",
        ]
    )
    _bold_header(ws)
    for r in nanopore_runs:
        ws.append(
            [
                r["run_accession"],
                r["sample_code"],
                r["sampling_date"],
                r["barcode"],
                r["minknow_sample_id"],
                r["alias"],
                r["protocol_id"],
                r["sequencing_kit_id"],
                r["type"],
                r["runName"],
                r["sampleName"],
                r["comments"],
            ]
        )

    ws = wb.create_sheet("biomeme")
    ws.append(
        [
            "biomeme_run_name",
            "sample_code",
            "sampling_date",
            "biomeme_sample_id",
            "dilution_factor",
            "comments",
        ]
    )
    _bold_header(ws)
    for r in biomeme_runs:
        ws.append(
            [
                r["biomeme_run_name"],
                r["sample_code"],
                r["sampling_date"],
                r["biomeme_sample_id"],
                r["dilution_factor"],
                r["comments"],
            ]
        )

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=odin_metadata.xlsx"},
    )


def _import_sites(
    db: sqlite3.Connection, wb: openpyxl.Workbook, now: str, actor: str,
    report: _ImportReport, seen_site_codes: set,
) -> None:
    if "sites" not in wb.sheetnames:
        return
    for r in _sheet_rows_by_header(wb["sites"], _AUTOFILLED_COLUMNS["sites"]):
        site_code = _str(r.get("site_code"))
        country = _str(r.get("country"))
        if not site_code or not country:
            missing = "site_code" if not site_code else "country"
            report.skip("sites", r["__row__"], f"{missing} is required and is empty")
            continue
        seen_site_codes.add(site_code)
        row = db.execute("SELECT * FROM sites WHERE site_code = ?", (site_code,)).fetchone()
        if row:
            all_fields = {
                "site": _str(r.get("site")),
                "country": country,
                "country_code": _str(r.get("country_code")),
                "city_code": _str(r.get("city_code")),
                "city": _str(r.get("city")),
                "location": _str(r.get("location")),
                "longitude": _flt(r.get("longitude")),
                "latitude": _flt(r.get("latitude")),
                "comments": _str(r.get("comments")),
            }
            # A column absent from the sheet says nothing about the field and leaves it
            # alone; without this, the real curated workbook — sites keyed by site_ID +
            # country only — would clear the components of every matched site. A blank
            # cell in a *present* column still clears, as everywhere in the importer.
            candidates = {key: value for key, value in all_fields.items() if key in r}

            # The API rules the importer used to bypass, applied with the importer's
            # skip-and-report contract instead of the endpoints' 409s:
            # country_code/city_code/site are frozen once samples reference the site
            # (their sample_code — half the sync merge key, and the name of pipeline
            # output on disk — is derived from them), and a short code must not be
            # re-bound to a different country/city name.
            frozen = [
                field
                for field in ("country_code", "city_code", "site")
                if field in candidates and candidates[field] != row[field]
            ]
            if frozen:
                sample_count = db.execute(
                    "SELECT COUNT(*) FROM samples WHERE site_id = ?", (row["id"],)
                ).fetchone()[0]
                if sample_count:
                    report.skip(
                        "sites", r["__row__"],
                        f"cannot change {', '.join(frozen)} on site '{site_code}': "
                        f"{sample_count} sample(s) reference it, and their sample_code is "
                        "derived from it. Create a new site instead.",
                    )
                    continue
            conflict = find_code_conflict(
                db,
                candidates.get("country_code", row["country_code"]),
                candidates.get("country", row["country"]),
                candidates.get("city_code", row["city_code"]),
                candidates.get("city", row["city"]),
                exclude_id=row["id"],
            )
            if conflict:
                report.skip("sites", r["__row__"], conflict)
                continue
            report.apply_update(db, "sites", row, candidates, now, actor, row=r["__row__"])
        else:
            conflict = find_code_conflict(
                db,
                _str(r.get("country_code")),
                country,
                _str(r.get("city_code")),
                _str(r.get("city")),
            )
            if conflict:
                report.skip("sites", r["__row__"], conflict)
                continue
            db.execute(
                """INSERT INTO sites
                   (id, site_code, site, country, country_code, city_code, city,
                    location, longitude, latitude, comments, created_at, updated_at, created_by, updated_by)
                   VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (
                    str(uuid.uuid4()),
                    site_code,
                    _str(r.get("site")),
                    country,
                    _str(r.get("country_code")),
                    _str(r.get("city_code")),
                    _str(r.get("city")),
                    _str(r.get("location")),
                    _flt(r.get("longitude")),
                    _flt(r.get("latitude")),
                    _str(r.get("comments")),
                    now,
                    now,
                    actor,
                    actor,
                ),
            )
            report.summary["sites"]["created"] += 1


def _import_samples(
    db: sqlite3.Connection, wb: openpyxl.Workbook, now: str, actor: str,
    report: _ImportReport, seen_sample_keys: set,
) -> None:
    if "samples" not in wb.sheetnames:
        return
    for r in _sheet_rows_by_header(wb["samples"], _AUTOFILLED_COLUMNS["samples"]):
        sample_code = _str(r.get("sample_code"))
        sampling_date = _str(r.get("sampling_date"))
        if not sample_code or not sampling_date:
            missing = "sample_code" if not sample_code else "sampling_date"
            report.skip("samples", r["__row__"], f"{missing} is required and is empty")
            continue
        # Marked as seen *before* the resolution check below, so that a row we skip still
        # protects an existing sample from replace-mode pruning. Skipping an unimportable row
        # must not delete good data.
        seen_sample_keys.add((sample_code, sampling_date))
        site_id = _find_site_id(db, _str(r.get("site_code")))
        sample_type = _str(r.get("sample_type"))
        date_extraction = _str(r.get("date_extraction"))
        # Validate the dates here rather than letting the schema CHECK reject them. The
        # constraint aborts the whole transaction, so a single malformed cell would fail the
        # entire import and roll back every good row with it — one bad date in a
        # hand-edited spreadsheet should cost that row, not the upload.
        if not is_yyyymmdd(sampling_date):
            report.skip(
                "samples", r["__row__"],
                f"sampling_date '{sampling_date}' is not YYYYMMDD",
            )
            continue
        if date_extraction is not None and not is_yyyymmdd(date_extraction):
            report.skip(
                "samples", r["__row__"],
                f"date_extraction '{date_extraction}' is not YYYYMMDD",
            )
            continue
        # sample_code is derived from site + sample_type and is half of the sync merge key,
        # so importing a row whose site_code does not resolve (or with no sample_type) would
        # store a sample whose code cannot be trusted. Skip it, as seed_samples already does,
        # rather than write a row that will collide or duplicate on the next merge.
        if site_id is None:
            report.skip(
                "samples", r["__row__"],
                f"site_code '{_str(r.get('site_code'))}' does not match any site "
                "(add it to the sites sheet, or correct the spelling)",
            )
            continue
        if not sample_type:
            report.skip("samples", r["__row__"], "sample_type is required and is empty")
            continue
        if not _lookup_code_exists(db, "sample_type", sample_type):
            report.skip(
                "samples", r["__row__"],
                f"sample_type '{sample_type}' is not a known lookup value",
            )
            continue
        row = db.execute(
            "SELECT * FROM samples WHERE sample_code = ? AND sampling_date = ?",
            (sample_code, sampling_date),
        ).fetchone()
        if row:
            all_fields = {
                "site_id": site_id,
                "sample_type": sample_type,
                "depth": _str(r.get("depth")),
                "elevation": _str(r.get("elevation")),
                "partner_sample_code": _str(r.get("partner_sample_code")),
                "nucleic_acid_concentration": _str(r.get("nucleic_acid_concentration")),
                "extract_volume": _str(r.get("extract_volume")),
                "elution_volume": _str(r.get("elution_volume")),
                "date_extraction": date_extraction,
                "comments_sampling": _str(r.get("comments_sampling")),
                "comments_extraction": _str(r.get("comments_extraction")),
                "comments": _str(r.get("comments")),
            }
            # An absent column leaves the field alone (see _import_sites). site_id and
            # sample_type are always present here: their source columns are required and
            # the row was skipped above when either was missing.
            candidates = {
                key: value
                for key, value in all_fields.items()
                if key in ("site_id", "sample_type") or key in r
            }
            report.apply_update(db, "samples", row, candidates, now, actor, row=r["__row__"])
        else:
            db.execute(
                """INSERT INTO samples
                   (id, sample_code, site_id, sample_type, depth, elevation, sampling_date,
                    partner_sample_code, nucleic_acid_concentration, extract_volume, elution_volume,
                    date_extraction, comments_sampling, comments_extraction, comments,
                    created_at, updated_at, created_by, updated_by)
                   VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (
                    str(uuid.uuid4()),
                    sample_code,
                    site_id,
                    sample_type,
                    _str(r.get("depth")),
                    _str(r.get("elevation")),
                    sampling_date,
                    _str(r.get("partner_sample_code")),
                    _str(r.get("nucleic_acid_concentration")),
                    _str(r.get("extract_volume")),
                    _str(r.get("elution_volume")),
                    date_extraction,
                    _str(r.get("comments_sampling")),
                    _str(r.get("comments_extraction")),
                    _str(r.get("comments")),
                    now,
                    now,
                    actor,
                    actor,
                ),
            )
            report.summary["samples"]["created"] += 1


def _import_nanopore(
    db: sqlite3.Connection, wb: openpyxl.Workbook, now: str, actor: str,
    report: _ImportReport, seen_nanopore_keys: set,
) -> None:
    if "nanopore" not in wb.sheetnames:
        return
    for r in _sheet_rows_by_header(wb["nanopore"], _AUTOFILLED_COLUMNS["nanopore"]):
        run_accession = _str(r.get("run_accession"))
        barcode = _str(r.get("barcode"))
        if not run_accession or not barcode:
            missing = "run_accession" if not run_accession else "barcode"
            report.skip("nanopore", r["__row__"], f"{missing} is required and is empty")
            continue
        seen_nanopore_keys.add((run_accession, barcode))

        # A sample reference that does not resolve is the same class of problem as an
        # unresolvable site_code on the samples sheet: the row would be stored pointing at
        # nothing. Only checked when a reference was actually given — a barcode may legitimately
        # be registered before its sample metadata exists.
        sample_code_ref = _str(r.get("sample_code"))
        if sample_code_ref is not None and resolve_sample_id(
            db, sample_code_ref, _str(r.get("sampling_date"))
        ) is None:
            report.skip(
                "nanopore", r["__row__"],
                f"sample_code '{sample_code_ref}' does not match any sample "
                "(check the sample_code and sampling_date)",
            )
            continue
        bad_lookup = _first_unknown_lookup(
            db,
            (
                ("protocol_id", _str(r.get("protocol_id"))),
                ("sequencing_kit_id", _str(r.get("sequencing_kit_id"))),
                ("mpox_type", _str(r.get("type"))),
            ),
        )
        if bad_lookup:
            field, value = bad_lookup
            report.skip(
                "nanopore", r["__row__"],
                f"{field} '{value}' is not a known lookup value",
            )
            continue

        accession_row = db.execute(
            "SELECT id FROM nanopore_run_accessions WHERE run_accession = ?",
            (run_accession,),
        ).fetchone()
        if accession_row:
            accession_id = accession_row["id"]
            # As everywhere in the importer: an absent column leaves the field alone, and
            # only values that actually differ are written. Absence is not the same as an
            # empty value, and updated_at drives sync's incoming_is_newer — so writing
            # unconditionally would clear untouched columns and bump updated_at on every
            # matched row, corrupting sync ordering.
            stored = db.execute(
                "SELECT * FROM nanopore_run_accessions WHERE id = ?", (accession_id,)
            ).fetchone()
            run_level = {
                key: _str(r.get(key))
                for key in ("protocol_id", "sequencing_kit_id", "runName", "sampleName", "comments")
                if key in r
            }
            changed = {key: value for key, value in run_level.items() if stored[key] != value}
            if changed:
                sql, params = build_update(
                    "nanopore_run_accessions", changed, accession_id, now, actor
                )
                db.execute(sql, params)
                report.record_change("nanopore", r["__row__"], stored, changed)
        else:
            accession_id = str(uuid.uuid4())
            db.execute(
                """INSERT INTO nanopore_run_accessions
                   (id, run_accession, protocol_id, sequencing_kit_id, runName, sampleName, comments,
                    created_at, updated_at, created_by, updated_by)
                   VALUES (?,?,?,?,?,?,?,?,?,?,?)""",
                (
                    accession_id,
                    run_accession,
                    _str(r.get("protocol_id")),
                    _str(r.get("sequencing_kit_id")),
                    _str(r.get("runName")),
                    _str(r.get("sampleName")),
                    _str(r.get("comments")),
                    now,
                    now,
                    actor,
                    actor,
                ),
            )

        sample_id = resolve_sample_id(
            db,
            _str(r.get("sample_code")),
            _str(r.get("sampling_date")),
        )
        nr = db.execute(
            "SELECT * FROM nanopore_runs WHERE accession_id = ? AND barcode = ?",
            (accession_id, barcode),
        ).fetchone()
        if nr:
            # sample_id only when the sheet has a sample_code column: an absent column must
            # not unlink the barcode from its sample. A present-but-blank cell still
            # unlinks, matching the API.
            candidates: dict[str, Any] = {}
            if "sample_code" in r:
                candidates["sample_id"] = sample_id
            if "type" in r:
                candidates["type"] = _str(r.get("type"))
            report.apply_update(db, "nanopore", nr, candidates, now, actor, row=r["__row__"])
        else:
            db.execute(
                """INSERT INTO nanopore_runs
                   (id, accession_id, sample_id, barcode, type, created_at, updated_at, created_by, updated_by)
                   VALUES (?,?,?,?,?,?,?,?,?)""",
                (
                    str(uuid.uuid4()),
                    accession_id,
                    sample_id,
                    barcode,
                    _str(r.get("type")),
                    now,
                    now,
                    actor,
                    actor,
                ),
            )
            report.summary["nanopore"]["created"] += 1


def _import_biomeme(
    db: sqlite3.Connection, wb: openpyxl.Workbook, now: str, actor: str,
    report: _ImportReport, seen_biomeme_names: set,
) -> None:
    if "biomeme" not in wb.sheetnames:
        return
    for r in _sheet_rows_by_header(wb["biomeme"], _AUTOFILLED_COLUMNS["biomeme"]):
        run_name = _str(r.get("biomeme_run_name"))
        if not run_name:
            report.skip("biomeme", r["__row__"], "biomeme_run_name is required and is empty")
            continue
        seen_biomeme_names.add(run_name)

        sample_code_ref = _str(r.get("sample_code"))
        sample_id = resolve_sample_id(db, sample_code_ref, _str(r.get("sampling_date")))
        # As on the nanopore sheet: a reference that was given must resolve. sample_id stays
        # nullable here because an unlinked biomeme run is a supported state.
        if sample_code_ref is not None and sample_id is None:
            report.skip(
                "biomeme", r["__row__"],
                f"sample_code '{sample_code_ref}' does not match any sample "
                "(check the sample_code and sampling_date)",
            )
            continue
        existing = db.execute(
            "SELECT * FROM biomeme_runs WHERE biomeme_run_name = ?",
            (run_name,),
        ).fetchone()
        if existing:
            all_fields = {
                "biomeme_sample_id": _str(r.get("biomeme_sample_id")),
                "dilution_factor": _flt(r.get("dilution_factor")),
                "comments": _str(r.get("comments")),
            }
            # An absent column leaves the field alone (see _import_sites); the sample link
            # in particular must survive a sheet with no sample_code column.
            candidates = {key: value for key, value in all_fields.items() if key in r}
            if "sample_code" in r:
                candidates["sample_id"] = sample_id
            report.apply_update(db, "biomeme", existing, candidates, now, actor, row=r["__row__"])
        else:
            db.execute(
                """INSERT INTO biomeme_runs
                   (id, biomeme_run_name, sample_id, biomeme_sample_id, dilution_factor, comments,
                    created_at, updated_at, created_by, updated_by)
                   VALUES (?,?,?,?,?,?,?,?,?,?)""",
                (
                    str(uuid.uuid4()),
                    run_name,
                    sample_id,
                    _str(r.get("biomeme_sample_id")),
                    _flt(r.get("dilution_factor")),
                    _str(r.get("comments")),
                    now,
                    now,
                    actor,
                    actor,
                ),
            )
            report.summary["biomeme"]["created"] += 1


def _prune_replaced(
    db: sqlite3.Connection, wb: openpyxl.Workbook, report: _ImportReport,
    seen_site_codes: set, seen_sample_keys: set,
    seen_nanopore_keys: set, seen_biomeme_names: set,
) -> None:
    """Replace-mode: delete rows absent from the imported workbook, one sheet at a
    time (only sheets present in the workbook are synchronised). Rows still
    referenced by another table are kept and counted as skipped."""
    if "biomeme" in wb.sheetnames:
        existing_bio = db.execute("SELECT id, biomeme_run_name FROM biomeme_runs").fetchall()
        delete_bio_ids = [r["id"] for r in existing_bio if r["biomeme_run_name"] not in seen_biomeme_names]
        if delete_bio_ids:
            db.executemany("DELETE FROM biomeme_runs WHERE id = ?", [(rid,) for rid in delete_bio_ids])
            report.summary["biomeme"]["deleted"] += len(delete_bio_ids)

    if "nanopore" in wb.sheetnames:
        existing_nano = db.execute(
            """SELECT nr.id, nra.run_accession, nr.barcode
               FROM nanopore_runs nr
               JOIN nanopore_run_accessions nra ON nra.id = nr.accession_id
               WHERE nra.run_accession IS NOT NULL"""
        ).fetchall()
        delete_nano_ids = [
            r["id"]
            for r in existing_nano
            if (r["run_accession"], r["barcode"]) not in seen_nanopore_keys
        ]
        if delete_nano_ids:
            db.executemany("DELETE FROM nanopore_runs WHERE id = ?", [(rid,) for rid in delete_nano_ids])
            report.summary["nanopore"]["deleted"] += len(delete_nano_ids)

        # Cleanup orphan accessions after nanopore row deletions.
        db.execute(
            """DELETE FROM nanopore_run_accessions
               WHERE run_accession IS NOT NULL
                 AND id NOT IN (SELECT DISTINCT accession_id FROM nanopore_runs)"""
        )

    if "samples" in wb.sheetnames:
        existing_samples = db.execute(
            "SELECT id, sample_code, sampling_date FROM samples"
        ).fetchall()
        candidate_delete_sample_ids = {
            r["id"]
            for r in existing_samples
            if (r["sample_code"], r["sampling_date"]) not in seen_sample_keys
        }
        if candidate_delete_sample_ids:
            # Keep referenced samples if nanopore/biomeme sheets are not part of this import.
            placeholders = ",".join(["?"] * len(candidate_delete_sample_ids))
            refs_nano = {
                r["sample_id"]
                for r in db.execute(
                    f"SELECT DISTINCT sample_id FROM nanopore_runs WHERE sample_id IN ({placeholders})",  # noqa: S608
                    tuple(candidate_delete_sample_ids),
                ).fetchall()
            }
            refs_bio = {
                r["sample_id"]
                for r in db.execute(
                    f"SELECT DISTINCT sample_id FROM biomeme_runs WHERE sample_id IN ({placeholders})",  # noqa: S608
                    tuple(candidate_delete_sample_ids),
                ).fetchall()
            }
            blocked_sample_ids = refs_nano | refs_bio
            delete_sample_ids = sorted(candidate_delete_sample_ids - blocked_sample_ids)
            if delete_sample_ids:
                db.executemany("DELETE FROM samples WHERE id = ?", [(rid,) for rid in delete_sample_ids])
                report.summary["samples"]["deleted"] += len(delete_sample_ids)
            if blocked_sample_ids:
                report.summary["samples"]["skipped"] += len(blocked_sample_ids)

    if "sites" in wb.sheetnames:
        existing_sites = db.execute("SELECT id, site_code FROM sites").fetchall()
        candidate_delete_site_ids = {
            r["id"]
            for r in existing_sites
            if r["site_code"] not in seen_site_codes
        }
        if candidate_delete_site_ids:
            placeholders = ",".join(["?"] * len(candidate_delete_site_ids))
            blocked_site_ids = {
                r["site_id"]
                for r in db.execute(
                    f"SELECT DISTINCT site_id FROM samples WHERE site_id IN ({placeholders})",  # noqa: S608
                    tuple(candidate_delete_site_ids),
                ).fetchall()
            }
            delete_site_ids = sorted(candidate_delete_site_ids - blocked_site_ids)
            if delete_site_ids:
                db.executemany("DELETE FROM sites WHERE id = ?", [(rid,) for rid in delete_site_ids])
                report.summary["sites"]["deleted"] += len(delete_site_ids)
            if blocked_site_ids:
                report.summary["sites"]["skipped"] += len(blocked_site_ids)


@router.post("/excel/import", summary="Import metadata from exported Excel")
async def import_excel(
    file: UploadFile = File(...),
    mode: Literal["merge", "replace"] = Query(default="merge"),
    dry_run: bool = Query(
        default=False,
        description=(
            "Validate and report what the upload would do, then roll back and write nothing. "
            "Returns 200 with the same summary and problem list an apply would produce, "
            "because it is the same code path — so the preview cannot disagree with the "
            "result. Never returns 422: a preview is an answer, not a failure."
        ),
    ),
    import_valid_rows: bool = Query(
        default=False,
        description=(
            "When false (the default) an upload containing any unimportable row is rejected "
            "with 422 and the full list of problems, and nothing is written. Set true to "
            "import the valid rows and skip the rest — intended to be set by the client only "
            "after the operator has seen that list and chosen to proceed."
        ),
    ),
    db: sqlite3.Connection = Depends(get_db),
):
    name = (file.filename or "").lower()
    if not name.endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="Please upload an .xlsx file")

    try:
        payload = await file.read()
        wb = openpyxl.load_workbook(io.BytesIO(payload), data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Invalid Excel file: {exc}") from exc

    now = _now()
    actor = "excel-import"
    replace_mode = mode == "replace"
    report = _ImportReport(
        summary={
            "sites": {"created": 0, "updated": 0, "unchanged": 0, "deleted": 0, "skipped": 0},
            "samples": {"created": 0, "updated": 0, "unchanged": 0, "deleted": 0, "skipped": 0},
            "nanopore": {"created": 0, "updated": 0, "unchanged": 0, "deleted": 0, "skipped": 0},
            "biomeme": {"created": 0, "updated": 0, "unchanged": 0, "deleted": 0, "skipped": 0},
        }
    )

    seen_site_codes: set[str] = set()
    seen_sample_keys: set[tuple[str, str]] = set()
    seen_nanopore_keys: set[tuple[str, str]] = set()
    seen_biomeme_names: set[str] = set()

    try:
        db.execute("BEGIN")

        _import_sites(db, wb, now, actor, report, seen_site_codes)
        _import_samples(db, wb, now, actor, report, seen_sample_keys)
        _import_nanopore(db, wb, now, actor, report, seen_nanopore_keys)
        _import_biomeme(db, wb, now, actor, report, seen_biomeme_names)

        if replace_mode:
            _prune_replaced(
                db, wb, report, seen_site_codes, seen_sample_keys,
                seen_nanopore_keys, seen_biomeme_names,
            )

        # Everything above ran inside the transaction, so the report is complete and nothing
        # is durable yet — which is what makes a preview possible without a second code path.
        # A preview is the real import, rolled back: whatever it reports is exactly what
        # applying would do, and the two cannot drift apart the way parallel implementations
        # would. Note a preview is never an error, however bad the workbook is; it is an
        # answer to "what would this do?".
        if dry_run:
            db.rollback()
            return {
                "detail": "Preview only — nothing was written.",
                "dry_run": True,
                "summary": report.summary,
                "problems": report.problems[:_MAX_REPORTED_PROBLEMS],
                "problem_count": len(report.problems),
                "problems_truncated": max(0, len(report.problems) - _MAX_REPORTED_PROBLEMS),
                "changes": report.changes[:_MAX_REPORTED_PROBLEMS],
                "change_count": len(report.changes),
                "changes_truncated": max(0, len(report.changes) - _MAX_REPORTED_PROBLEMS),
                "mode": mode,
            }

        # Default to rejecting the whole upload and telling the operator exactly what to fix;
        # import the valid rows only when explicitly asked.
        if report.problems and not import_valid_rows:
            raise HTTPException(
                status_code=422,
                detail=report.as_detail(
                    f"{len(report.problems)} row(s) cannot be imported, so nothing was "
                    "imported. Fix the rows listed below and upload again, or re-send with "
                    "import_valid_rows=true to import only the valid rows."
                ),
            )

        db.commit()
    except sqlite3.IntegrityError as exc:
        db.rollback()
        raise HTTPException(status_code=409, detail=f"Import failed due to constraint conflict: {exc}") from exc
    except HTTPException:
        db.rollback()
        raise
    except Exception as exc:
        db.rollback()
        raise HTTPException(status_code=400, detail=f"Import failed: {exc}") from exc

    detail = "Import completed"
    if replace_mode:
        detail = "Import completed (replace mode)"
    return {
        "detail": detail,
        "summary": report.summary,
        "problems": report.problems[:_MAX_REPORTED_PROBLEMS],
        "problem_count": len(report.problems),
        "mode": mode,
    }
